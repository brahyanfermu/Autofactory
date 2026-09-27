"""
Módulo para exportar datos a Excel usando openpyxl.
"""
import datetime
from tkinter import messagebox, filedialog
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def exportar_a_excel(datos, columnas, titulo, nombre_archivo):
    """
    Exporta una lista de diccionarios a un archivo Excel.

    Parámetros:
        datos: lista de diccionarios con los datos
        columnas: lista de tuplas (clave, título_mostrar)
        titulo: título de la hoja (ej: "Vehículos")
        nombre_archivo: nombre sugerido del archivo (sin extensión)
    """
    if not datos:
        messagebox.showwarning("Sin datos", "No hay datos para exportar.")
        return

    # Preguntar dónde guardar
    ruta = filedialog.asksaveasfilename(
        title="Guardar como Excel",
        defaultextension=".xlsx",
        initialfile=f"{nombre_archivo}_{datetime.date.today()}.xlsx",
        filetypes=[("Archivos Excel", "*.xlsx")]
    )

    if not ruta:
        return  # Usuario canceló

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = titulo[:31]  # Excel limita a 31 caracteres

        # ---------- Estilos ----------
        color_header = "1B2A41"
        color_texto_header = "FFFFFF"

        font_header = Font(bold=True, color=color_texto_header, size=12)
        fill_header = PatternFill("solid", fgColor=color_header)
        alineacion_centro = Alignment(horizontal="center", vertical="center")
        borde_fino = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

        # ---------- Encabezados ----------
        for col_idx, (clave, header) in enumerate(columnas, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = alineacion_centro
            cell.border = borde_fino

        # ---------- Datos ----------
        for row_idx, item in enumerate(datos, start=2):
            for col_idx, (clave, _) in enumerate(columnas, start=1):
                valor = item.get(clave, "")
                cell = ws.cell(row=row_idx, column=col_idx, value=valor)
                cell.border = borde_fino
                cell.alignment = Alignment(horizontal="center")

        # ---------- Ancho automático de columnas ----------
        for col_idx, (clave, header) in enumerate(columnas, start=1):
            max_len = len(header)
            for item in datos:
                valor = str(item.get(clave, ""))
                max_len = max(max_len, len(valor))
            ws.column_dimensions[
                ws.cell(row=1, column=col_idx).column_letter
            ].width = max_len + 4

        wb.save(ruta)
        messagebox.showinfo("Éxito",
            f"Datos exportados correctamente:\n\n{ruta}")
    except Exception as e:
        messagebox.showerror("Error", f"No se pudo exportar:\n\n{e}")